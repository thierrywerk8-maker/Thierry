"""Rebuild the 'Door2Door| Paylogic' planning for di 13 / wo 14 / do 15 oktober.

Usage: python3 -I build_planning_v2.py <input.xlsx> <output.xlsx>
Input is Thierry's edited workbook (tabs 'Door2Door| Paylogic', 'Door2Door | Eventim (Light)',
'Leadlijst sheet. '). Team per lead comes from the right-hand table of the Leadlijst.
"""
import sys, datetime
from urllib.parse import quote
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
PLAN, LIJST = 'Door2Door| Paylogic', 'Leadlijst sheet. '

# ---- lead data: address/place/link from the old planning, owner + team from the Leadlijst ----
leads = {}
old = wb[PLAN]
for r in range(1, old.max_row + 1):
    lead, stop = old.cell(r, 3).value, old.cell(r, 2).value
    if isinstance(stop, int) and lead and lead not in leads:
        k = old.cell(r, 11)
        leads[lead] = dict(addr=old.cell(r, 5).value, place=old.cell(r, 6).value,
                           link=k.hyperlink.target if k.hyperlink else None)
lijst = wb[LIJST]
team_of, LIJST_ROWS = {}, []
for r in range(6, lijst.max_row + 1):
    lead = lijst.cell(r, 13).value  # kolom M
    if lead:
        leads[lead]['owner'] = lijst.cell(r, 14).value  # kolom N
        team_of[lead] = lijst.cell(r, 15).value          # kolom O
        LIJST_ROWS.append(r)

MOVED = {'Paard', 'Audio Obscura', 'Gashouder', 'Dfrnt Festival (Tapt)', 'Lofi'}  # donkergroen in de Leadlijst
NOT_PLANNED = {'Audio Obscura': 'Past evt. bij Friso + Thierry op donderdag (terug dan ca. 15:15)',
               'Paard': 'Past evt. bij Friso + Thierry op woensdag, Rotterdam-dag (terug dan ca. 17:15)'}
MEMBERS = {'Y': {'Yasmijn', 'Steven'}, 'F': {'Friso', 'Thierry'}}
LABEL = {('Y', 1): 'Yasmijn + Steven + Milo', ('Y', 2): 'Yasmijn + Steven', ('Y', 3): 'Yasmijn + Steven',
         ('F', 1): 'Friso + Thierry', ('F', 2): 'Friso + Thierry + Cecil', ('F', 3): 'Friso + Thierry'}
DAYNAME = {1: 'DINSDAG 13 OKTOBER', 2: 'WOENSDAG 14 OKTOBER', 3: 'DONDERDAG 15 OKTOBER'}
KANTOOR = 'KANTOOR'

# (lead, rijtijd vanaf vorige stop in min); laatste getal = rijtijd terug naar AC
ROUTES = {
    ('Y', 1): ([('Intercell/ AFP', 11), ('Jaarbeurs', 39), ('Elevation Events', 5), ('Studio Linksvoor', 63),
                ('Classified Events [Urgent]', 14), ('Listen to Flow', 30), ('ONA Groep', 61), ('Reggae Lake', 45)], 12),
    ('Y', 2): ([('Gashouder', 14), ('Dfrnt Festival (Tapt)', 13), ('Lofi', 10), ('Kaleidoscope', 126),
                ('Paradigm', 17), ('Sunsplash', 124)], 26),
    ('Y', 3): ([('Muze Misse', 75), ('Par-t', 25), ('Minority Events', 40), ('Volt', 67)], 132),
    ('F', 1): ([('Making Waves', 20), ('VBX', 8), ('Thuishaven', 12), ('IsBurning', 10), ('Chasing the Hihat', 12),
                (KANTOOR, 13),
                ('Studio Wieman', 12), ('ingang / no clubs', 8), ('A Venue', 18), ('Melkweg', 5)], 14),
    ('F', 2): ([('Cinema Culinair', 70), ('Give Soul', 10), ('CUBE', 7), ('Ahoy Rotterdam', 12), ('The Zoo', 18),
                ('Vivre Productions', 27), ('White House Scheveningen', 49)], 93),
    ('F', 3): ([('8Ball', 38), ('Verknipt', 31), ('3X3 Basketbal', 17), ('Loveland', 22), ('Vault Sessions', 23),
                ('Tropikali', 27)], 14),
}
planned = [l for stops, _ in ROUTES.values() for l, _ in stops if l != KANTOOR]
assert len(planned) == len(set(planned)), 'lead twice in the planning'
expected = set(team_of) - set(NOT_PLANNED)
assert set(planned) == expected, (set(planned) ^ expected)
for (t, d), (stops, _) in ROUTES.items():
    for l, _ in stops:
        if l != KANTOOR:
            assert team_of[l].startswith('Yasmijn' if t == 'Y' else 'Friso'), (l, t)

def maps_link(q):
    return 'https://www.google.com/maps/search/?api=1&query=' + quote(q, safe='/')

# ---- styles (same as the existing planning) ----
WHITE_B14 = Font(name='Calibri', size=14, bold=True, color='FFFFFFFF')
WHITE_B12 = Font(name='Calibri', size=12, bold=True, color='FFFFFFFF')
WHITE_B11 = Font(name='Calibri', size=11, bold=True, color='FFFFFFFF')
BOLD = Font(name='Calibri', size=11, bold=True)
INPUT = Font(name='Calibri', size=11, bold=True, color='FF0000FF')
LINK = Font(name='Calibri', size=11, color='FF087E68', underline='single')
TITLE_FILL = PatternFill('solid', fgColor='FF17243A')
HEAD_FILL = PatternFill('solid', fgColor='FF244466')
ENDS_FILL = PatternFill('solid', fgColor='FFE8EFF5')
TAKEN_FILL = PatternFill('solid', fgColor='FFFFF0CE')
NEW_FILL = PatternFill('solid', fgColor='FFE2EFDA')
OFFICE_FILL = PatternFill('solid', fgColor='FFDDEBF7')
RED_FILL = PatternFill('solid', fgColor='FFFFC7CE')
GREEN_FILL = PatternFill('solid', fgColor='FFC6EFCE')
WRAP = Alignment(wrap_text=True, vertical='top')
HHMM = 'h:mm'
NCOL = 11
RED_RULE = dict(fill=RED_FILL, font=Font(bold=True, color='FF9C0006'))

del wb[PLAN]
ws = wb.create_sheet(PLAN, 0)
ws.sheet_properties.tabColor = '244466'
for col, w in zip('ABCDEFGHIJK', [24, 12, 36, 15, 46, 24, 19, 14, 14, 46, 18]):
    ws.column_dimensions[col].width = w

def bar(row, text, font, fill):
    ws.cell(row, 1, text)
    for c in range(1, NCOL + 1):
        ws.cell(row, c).font, ws.cell(row, c).fill = font, fill

def header(row, labels):
    for c, v in enumerate(labels, 1):
        cell = ws.cell(row, c, v)
        cell.font, cell.fill = WHITE_B11, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')

def late(limit_h, limit_m):
    return CellIsRule(operator='greaterThan', formula=[f'TIME({limit_h},{limit_m},0)'], **RED_RULE)

# ---- titel + uitgangspunten ----
bar(1, 'ADE Door-to-Door | Paylogic | 2 teams x 3 dagen (di 13 t/m do 15 oktober)', WHITE_B14, TITLE_FILL)
ws.row_dimensions[1].height = 18.95
info = [
    ('Start / eind', 'Amsterdam Centraal', 'Start 09:30', 'Uiterlijk circa 17:30'),
    ('Bezoekduur', '20 minuten', 'Lunch', '30 minuten, ingepland'),
    ('Routeplanning', 'Indicatief: rijtijden handmatig geschat; geen live Maps-tijden'),
    ('Let op', 'Dinsdag zit Friso van 12:45 tot 14:45 op kantoor (De Ruijterkade 113); de pauze van Friso + Thierry valt '
               'in dat blok. Milo rijdt dinsdag mee met Yasmijn + Steven, Cecil woensdag met Friso + Thierry (Ahoy-dag).'),
]
for i, vals in enumerate(info, 2):
    for c, v in enumerate(vals, 1):
        ws.cell(i, c, v)
ws.cell(6, 1, 'Legenda')
for c, (txt, fill, font) in enumerate([('Verschoven naar Yasmijn + Steven (was Friso + Thierry)', NEW_FILL, None),
                                       ('Overgenomen van andere leadowner', TAKEN_FILL, None),
                                       ('Te laat (na 17:30 terug / na 12:45 op kantoor)', RED_FILL, None),
                                       ('Blauwe tekst = invoer (vertrek kantoor)', None, INPUT)], 3):
    cell = ws.cell(6, c, txt)
    if fill: cell.fill = fill
    if font: cell.font = font

OV_HEAD = 8
OV_FIRST = OV_HEAD + 1
ORDER = [('Y', d) for d in (1, 2, 3)] + [('F', d) for d in (1, 2, 3)]
row = OV_FIRST + len(ORDER) + 2

# ---- dagblokken ----
blocks, rows_of = {}, {}
for day in (1, 2, 3):
    bar(row, f'DAG {day}  |  {DAYNAME[day]}  |  VERTREK 09:30  |  20 MINUTEN PER LEAD + 30 MIN LUNCH', WHITE_B12, TITLE_FILL)
    ws.row_dimensions[row].height = 15.95
    row += 1
    header(row, ['Team', 'Stop', 'Lead', 'Leadowner', 'Adres', 'Plaats', 'Rijtijd vanaf vorig (min)',
                 'Aankomst', 'Vertrek', 'Opmerking', 'Google Maps'])
    row += 1
    for t in ('Y', 'F'):
        team = LABEL[(t, day)]
        stops, back = ROUTES[(t, day)]
        office_day = any(l == KANTOOR for l, _ in stops)
        n_leads = sum(1 for l, _ in stops if l != KANTOOR)
        start_note = 'Per auto; rijtijden indicatief'
        if '+ Milo' in team: start_note += ' | Milo rijdt mee'
        if '+ Cecil' in team: start_note += ' | Cecil rijdt mee'
        vals = [team, 'START', 'Amsterdam Centraal', None, 'Stationsplein, Amsterdam', 'Amsterdam', 0,
                datetime.time(9, 30), datetime.time(9, 30), start_note, None]
        for c, v in enumerate(vals, 1):
            ws.cell(row, c, v).fill = ENDS_FILL
        ws.cell(row, 8).number_format = ws.cell(row, 9).number_format = HHMM
        row += 1
        first, n = row, 0
        for lead, drive in stops:
            if lead == KANTOOR:
                vals = [team, KANTOOR, 'Paylogic kantoor | Friso afzetten en ophalen', 'Friso', 'De Ruijterkade 113',
                        'Amsterdam', drive, f'=I{row-1}+G{row}/1440', datetime.time(14, 45),
                        'Meeting Friso 12:45-14:45, pauze in dit blok. Uitloop? Pas vertrektijd (blauw) aan.', 'Open adres']
                for c, v in enumerate(vals, 1):
                    cell = ws.cell(row, c, v)
                    cell.fill, cell.alignment = OFFICE_FILL, WRAP
                ws.cell(row, 9).font = INPUT
                ws.cell(row, 11).hyperlink = maps_link('De Ruijterkade 113, Amsterdam Netherlands')
                ws.cell(row, 11).font = LINK
                ws.cell(row, 8).number_format = ws.cell(row, 9).number_format = HHMM
                ws.conditional_formatting.add(f'H{row}', late(12, 45))
                rows_of[KANTOOR] = row
                ws.row_dimensions[row].height = 33.95
                row += 1
                continue
            n += 1
            d = leads[lead]
            lunch = 30 if (n == 5 and not office_day) else 0
            own = d['owner'] in MEMBERS[t]
            note = 'Eigen lead' if own else f"Overgenomen van {d['owner']}"
            if lead in MOVED:
                note += ' (was Friso + Thierry)'
            vals = [team, n, lead, d['owner'], d['addr'], d['place'], drive,
                    f'=I{row-1}+G{row}/1440+{lunch}/1440', f'=H{row}+20/1440', note, 'Open adres']
            for c, v in enumerate(vals, 1):
                ws.cell(row, c, v).alignment = WRAP
            ws.cell(row, 8).number_format = ws.cell(row, 9).number_format = HHMM
            if not own:
                ws.cell(row, 10).fill = TAKEN_FILL
            if lead in MOVED:
                for c in (1, 2, 3):
                    ws.cell(row, c).fill = NEW_FILL
                ws.cell(row, 3).font = BOLD
            k = ws.cell(row, 11)
            k.hyperlink = d['link'] or maps_link(f"{d['addr'].replace(' | ', ', ')}, {d['place']} Netherlands")
            k.font = LINK
            ws.row_dimensions[row].height = 33.95
            rows_of[lead] = row
            row += 1
        last = row - 1
        lunch = 30 if (n_leads < 5 and not office_day) else 0
        lunch_note = 'Pauze tijdens meeting Friso' if office_day else 'Inclusief lunch van 30 min'
        vals = [team, 'TERUG', 'Amsterdam Centraal', None, 'Stationsplein', 'Amsterdam', back,
                f'=I{row-1}+G{row}/1440+{lunch}/1440', f'=H{row}',
                f'=IF(H{row}>TIME(17,30,0),"{lunch_note} | LET OP: na 17:30 terug, zie Aandachtspunten","{lunch_note}")', None]
        for c, v in enumerate(vals, 1):
            ws.cell(row, c, v).fill = ENDS_FILL
        ws.cell(row, 8).number_format = ws.cell(row, 9).number_format = HHMM
        ws.cell(row, 10).alignment = WRAP
        ws.conditional_formatting.add(f'H{row}:I{row}', late(17, 30))
        blocks[(t, day)] = (first, last, row, 0 if office_day else 30)
        row += 2
    row += 1
ROUTES_END = row - 1

# ---- overzicht ----
header(OV_HEAD, ['Team', 'Dag', 'Leads', 'Eigen leads', 'Overgenomen', 'Rijtijd (min)',
                 'Bezoektijd (min)', 'Lunch (min)', 'Terug ca.', 'Binnen 17:30?'])
for i, (t, day) in enumerate(ORDER):
    r = OV_FIRST + i
    first, last, back, lunch = blocks[(t, day)]
    vals = [LABEL[(t, day)], DAYNAME[day].split()[0].capitalize(), f'=COUNT(B{first}:B{last})',
            f'=COUNTIF(J{first}:J{last},"Eigen lead*")', f'=COUNTIF(J{first}:J{last},"Overgenomen*")',
            f'=SUM(G{first}:G{back})', f'=C{r}*20', lunch, f'=H{back}', f'=IF(I{r}<=TIME(17,30,0),"JA","NEE")']
    for c, v in enumerate(vals, 1):
        ws.cell(r, c, v)
    ws.cell(r, 9).number_format = HHMM
    ws.row_dimensions[r].height = 24
tot = OV_FIRST + len(ORDER)
ws.cell(tot, 1, 'Totaal').font = BOLD
for c in 'CDEFG':
    ws[f'{c}{tot}'] = f'=SUM({c}{OV_FIRST}:{c}{tot-1})'
    ws[f'{c}{tot}'].font = BOLD
rng = f'J{OV_FIRST}:J{tot-1}'
ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"NEE"'], **RED_RULE))
ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"JA"'], fill=GREEN_FILL, font=Font(color='FF006100')))
ws.conditional_formatting.add(f'I{OV_FIRST}:I{tot-1}', late(17, 30))

# ---- verschoven naar Yasmijn + Steven ----
def lookup_day(ref, sheet=''):
    p = sheet
    match = f'MATCH({ref},{p}$C$1:$C${ROUTES_END},0)'
    return (f'=IFERROR(CHOOSE(COUNTIF({p}$A$1:INDEX({p}$A$1:$A${ROUTES_END},{match}),"DAG *"),'
            f'"Dinsdag","Woensdag","Donderdag"),"Niet ingepland")')

bar(row, 'VERSCHOVEN NAAR YASMIJN + STEVEN | donkergroen in de Leadlijst', WHITE_B12, TITLE_FILL)
row += 1
header(row, ['Nieuw team', 'Dag', 'Lead', 'Leadowner', 'Status', 'Plaats', None, 'Aankomst', None, 'Toelichting'])
row += 1
for lead in ['Gashouder', 'Dfrnt Festival (Tapt)', 'Lofi', 'Audio Obscura', 'Paard']:
    d = leads[lead]
    match = f'MATCH(C{row},$C$1:$C${ROUTES_END},0)'
    vals = [team_of[lead], lookup_day(f'C{row}'), lead, d['owner'],
            f'=IF(ISNUMBER({match}),"Ingepland","Niet ingepland | Yasmijn plant zelf in")', d['place'], None,
            f'=IFERROR(INDEX($H$1:$H${ROUTES_END},{match}),"")', None,
            NOT_PLANNED.get(lead, 'Eigen lead van Yasmijn')]
    for c, v in enumerate(vals, 1):
        ws.cell(row, c, v).alignment = WRAP
    ws.cell(row, 8).number_format = HHMM
    ws.cell(row, 3).fill = NEW_FILL
    if lead in NOT_PLANNED:
        ws.cell(row, 5).fill = RED_FILL
    row += 1
row += 1

# ---- aandachtspunten ----
bar(row, 'AANDACHTSPUNTEN', WHITE_B12, TITLE_FILL)
row += 1
notes = [
    ('Audio Obscura + Paard', 'Niet ingepland: Yasmijn kijkt zelf of ze deze op een andere dag kan doen. Als Friso + Thierry ze '
                              'toch meenemen: Audio Obscura past donderdag (terug ca. 15:15), Paard past woensdag op de '
                              'Rotterdam-dag (terug ca. 17:15).'),
    ('Friso + Thierry, dinsdag', 'Friso is om ca. 12:25 op kantoor (uiterlijk 12:45). Vertrek kantoor staat op 14:45 (blauwe cel): '
                                 'loopt de meeting uit, pas die tijd aan en de middag schuift mee. Bij 30 min uitloop '
                                 'terug ca. 17:32. Lukt een lead niet, dan kan die donderdag nog (die dag is om ca. 14:52 klaar).'),
    ('Yasmijn + Steven, woensdag', 'Terug ca. 17:30, geen speling: twee keer ca. 2 uur rijden (Amsterdam-Groningen en terug '
                                   'via Zuidoost).'),
    ('Classified Events', 'Staat als [Urgent] op dinsdag, de eerste dag.'),
    ('Volt', 'Leadowner Thierry (dubbel met Poppodium Volt), bezoek door Yasmijn + Steven op donderdag.'),
    ('Woodstock', 'Vervalt; niet meer ingepland.'),
    ('The Zoo', "Postcode staat als '274 LH'. Waarschijnlijk is dat 3274 LH (Heinenoord); graag controleren."),
    ('Rijtijden', 'Trajecten uit de eerdere planning houden hun rijtijd. Nieuwe trajecten zijn geschat (geen live Maps); de '
                  'heenrit naar Rotterdam is in de ochtendspits op 70 min gezet. Parkeren en spits kunnen afwijken.'),
]
for label, text in notes:
    ws.cell(row, 1, label).font = BOLD
    ws.cell(row, 1).alignment = WRAP
    ws.cell(row, 2, text)
    row += 1

# ---- Leadlijst rechter tabel: Dag (kolom P) uit de planning ----
P = f"'{PLAN}'!"
for r in LIJST_ROWS:
    lijst.cell(r, 16).value = lookup_day(f'M{r}', P)

wb.active = 0
for s in wb.worksheets:
    s.sheet_view.tabSelected = (s.title == PLAN)
wb.save(dst)
print('saved', dst, '| routes end row', ROUTES_END, '| blocks', blocks)

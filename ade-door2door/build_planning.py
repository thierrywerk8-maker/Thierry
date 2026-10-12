import sys, datetime
from urllib.parse import quote
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule, FormulaRule

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)

# ---- lead data from the existing day sheets ----
leads = {}
for name in ['Dag 1', 'Dag 2', 'Dag 3']:
    ws = wb[name]
    for r in range(3, ws.max_row + 1):
        team, stop, lead = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if not lead or stop in ('START', 'TERUG'):
            continue
        link = ws.cell(r, 11).hyperlink.target if ws.cell(r, 11).hyperlink else None
        leads[lead] = dict(owner=ws.cell(r, 4).value, addr=ws.cell(r, 5).value.strip(),
                           place=ws.cell(r, 6).value, link=link, team=team, day=int(name[-1]))
# normalise the two addresses that were typed differently in the source
leads['Listen to Flow']['addr'] = 'Roodwilligen 63 | 6903 TB'
leads['Jaarbeurs']['addr'] = 'Jaarbeursplein 6 | 3521 AL'
# Volt (dubbel met Poppodium Volt) valt nu onder Thierry
leads['Volt']['owner'] = 'Thierry'

def maps_link(addr, place):
    q = addr.replace(' | ', ', ') + ', ' + place + ' Netherlands'
    return 'https://www.google.com/maps/search/?api=1&query=' + quote(q, safe='/')

TEAMS = {'Yasmijn + Steven': {'Yasmijn', 'Steven'}, 'Friso + Thierry': {'Friso', 'Thierry'}}

# (lead, rijtijd vanaf vorig) ... and rijtijd terug naar AC
ROUTES = {
    ('Yasmijn + Steven', 1): ([('Paradigm', 130), ('Kaleidoscope', 17), ('ONA Groep', 115), ('Reggae Lake', 45)], 12),
    ('Yasmijn + Steven', 2): ([('Listen to Flow', 100), ('Classified Events [Urgent]', 30), ('Studio Linksvoor', 14),
                               ('Muze Misse', 35), ('Par-t', 25), ('Minority Events', 40)], 112),
    ('Yasmijn + Steven', 3): ([('Volt', 132), ('Elevation Events', 97), ('Jaarbeurs', 5), ('Intercell/ AFP', 39),
                               ('Sunsplash', 29)], 26),
    ('Friso + Thierry', 1): ([('White House Scheveningen', 93), ('Paard', 19), ('Cinema Culinair', 30), ('Give Soul', 10),
                              ('CUBE', 7), ('Ahoy Rotterdam', 12), ('The Zoo', 18), ('Vivre Productions', 27)], 79),
    ('Friso + Thierry', 2): ([('Studio Wieman', 15), ('ingang / no clubs', 8), ('Melkweg', 20), ('A Venue', 5),
                              ('Tropikali', 5), ('Audio Obscura', 4), ('Loveland', 25), ('3X3 Basketbal', 22),
                              ('Verknipt', 17), ('8Ball', 31)], 38),
    ('Friso + Thierry', 3): ([('Chasing the Hihat', 20), ('Gashouder', 8), ('Thuishaven', 9), ('IsBurning', 10),
                              ('Dfrnt Festival (Tapt)', 10), ('Lofi', 10), ('VBX', 10), ('Making Waves', 8),
                              ('Vault Sessions', 20), ('Woodstock', 30)], 44),
}
NEW = {l for l, d in leads.items() if d['team'] == 'Emiel + Bas' and d['owner'] not in ('Bas', 'Emiel')}
assert len(NEW) == 12, NEW
planned = [l for (stops, _) in ROUTES.values() for l, _ in stops]
assert sorted(planned) == sorted(l for l, d in leads.items() if d['team'] != 'Emiel + Bas' or l in NEW)

# ---- styles (taken from the existing sheets) ----
WHITE_B14 = Font(name='Calibri', size=14, bold=True, color='FFFFFFFF')
WHITE_B12 = Font(name='Calibri', size=12, bold=True, color='FFFFFFFF')
WHITE_B11 = Font(name='Calibri', size=11, bold=True, color='FFFFFFFF')
BOLD = Font(name='Calibri', size=11, bold=True)
LINK = Font(name='Calibri', size=11, color='FF087E68', underline='single')
TITLE_FILL = PatternFill('solid', fgColor='FF17243A')
HEAD_FILL = PatternFill('solid', fgColor='FF244466')
ENDS_FILL = PatternFill('solid', fgColor='FFE8EFF5')
TAKEN_FILL = PatternFill('solid', fgColor='FFFFF0CE')
NEW_FILL = PatternFill('solid', fgColor='FFE2EFDA')
RED_FILL = PatternFill('solid', fgColor='FFFFC7CE')
GREEN_FILL = PatternFill('solid', fgColor='FFC6EFCE')
WRAP = Alignment(wrap_text=True, vertical='top')
HHMM = 'h:mm'
NCOL = 11

ws = wb.create_sheet('Planning 2 teams', 0)
ws.sheet_properties.tabColor = '244466'
for col, w in zip('ABCDEFGHIJK', [22, 12, 36, 15, 46, 24, 19, 14, 14, 44, 18]):
    ws.column_dimensions[col].width = w

def bar(row, text, font, fill, ncol=NCOL):
    ws.cell(row, 1, text)
    for c in range(1, ncol + 1):
        ws.cell(row, c).font = font
        ws.cell(row, c).fill = fill

def header(row, labels):
    for c, v in enumerate(labels, 1):
        cell = ws.cell(row, c, v)
        cell.font, cell.fill = WHITE_B11, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')

late = CellIsRule(operator='greaterThan', formula=['TIME(17,30,0)'], fill=RED_FILL, font=Font(bold=True, color='FF9C0006'))

# ---- title + uitgangspunten ----
bar(1, 'ADE Door-to-Door | 2 teams x 3 dagen | herverdeling zonder Emiel + Bas', WHITE_B14, TITLE_FILL)
ws.row_dimensions[1].height = 18.95
info = [
    ('Start / eind', 'Amsterdam Centraal', 'Start 09:30', 'Uiterlijk circa 17:30'),
    ('Bezoekduur', '20 minuten', 'Lunch', '30 minuten, ingepland'),
    ('Routeplanning', 'Indicatief: rijtijden handmatig geschat; geen live Maps-tijden'),
    ('Let op', 'Emiel + Bas doen alleen hun eigen leads (zie Sheet1) en staan niet in deze planning. '
               'Hun overgenomen leads van andere leadowners zijn verdeeld over de twee teams hieronder.'),
]
for i, vals in enumerate(info, 2):
    for c, v in enumerate(vals, 1):
        ws.cell(i, c, v)
ws.cell(6, 1, 'Legenda')
for c, (txt, fill) in enumerate([('Nieuw toegevoegd (was Emiel + Bas)', NEW_FILL),
                                 ('Overgenomen van andere leadowner', TAKEN_FILL),
                                 ('Na 17:30 terug: oplossing nodig', RED_FILL)], 3):
    ws.cell(6, c, txt).fill = fill

OV_HEAD = 8
OV_FIRST = OV_HEAD + 1
order = [(t, d) for t in TEAMS for d in (1, 2, 3)]
row = OV_FIRST + len(order) + 2

# ---- day blocks ----
blocks = {}
new_rows = {}
for day in (1, 2, 3):
    bar(row, f'DAG {day}  |  VERTREK 09:30  |  20 MINUTEN PER LEAD + 30 MIN LUNCH', WHITE_B12, TITLE_FILL)
    ws.row_dimensions[row].height = 15.95
    row += 1
    header(row, ['Team', 'Stop', 'Lead', 'Leadowner', 'Adres', 'Plaats', 'Rijtijd vanaf vorig (min)',
                 'Aankomst', 'Vertrek', 'Opmerking', 'Google Maps'])
    row += 1
    for team in TEAMS:
        stops, back = ROUTES[(team, day)]
        start = row
        vals = [team, 'START', 'Amsterdam Centraal', None, 'Stationsplein, Amsterdam', 'Amsterdam', 0,
                datetime.time(9, 30), datetime.time(9, 30), 'Per auto; rijtijden indicatief', None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c, v)
            cell.fill = ENDS_FILL
        ws.cell(row, 8).number_format = ws.cell(row, 9).number_format = HHMM
        row += 1
        first = row
        lunch_on_return = len(stops) < 5
        for n, (lead, drive) in enumerate(stops, 1):
            d = leads[lead]
            lunch = 30 if n == 5 else 0
            own = d['owner'] in TEAMS[team]
            note = 'Eigen lead' if own else f"Overgenomen van {d['owner']}"
            if lead in NEW:
                note += f" (was Emiel + Bas, dag {d['day']})"
            vals = [team, n, lead, d['owner'], d['addr'], d['place'], drive,
                    f'=I{row-1}+G{row}/1440+{lunch}/1440', f'=H{row}+20/1440', note, 'Open adres']
            for c, v in enumerate(vals, 1):
                ws.cell(row, c, v).alignment = WRAP
            ws.cell(row, 8).number_format = ws.cell(row, 9).number_format = HHMM
            if not own:
                ws.cell(row, 10).fill = TAKEN_FILL
            if lead in NEW:
                for c in (1, 2, 3):
                    ws.cell(row, c).fill = NEW_FILL
                ws.cell(row, 3).font = BOLD
                new_rows[lead] = (team, day, row)
            k = ws.cell(row, 11)
            k.hyperlink = d['link'] or maps_link(d['addr'], d['place'])
            k.font = LINK
            ws.row_dimensions[row].height = 33.95
            row += 1
        last = row - 1
        lunch = 30 if lunch_on_return else 0
        vals = [team, 'TERUG', 'Amsterdam Centraal', None, 'Stationsplein', 'Amsterdam', back,
                f'=I{row-1}+G{row}/1440+{lunch}/1440', f'=H{row}',
                f'=IF(H{row}>TIME(17,30,0),"Inclusief lunch van 30 min | LET OP: na 17:30 terug, zie Aandachtspunten",'
                f'"Inclusief lunch van 30 min")', None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c, v)
            cell.fill = ENDS_FILL
        ws.cell(row, 8).number_format = ws.cell(row, 9).number_format = HHMM
        ws.cell(row, 10).alignment = WRAP
        ws.conditional_formatting.add(f'H{row}:I{row}', late)
        blocks[(team, day)] = (first, last, row)
        row += 2
    row += 1

# ---- overzicht (formulas pointing at the blocks above) ----
header(OV_HEAD, ['Team', 'Dag', 'Leads', 'Eigen leads', 'Nieuw toegevoegd', 'Rijtijd (min)',
                 'Bezoektijd (min)', 'Lunch (min)', 'Terug ca.', 'Binnen 17:30?'])
for i, (team, day) in enumerate(order):
    r = OV_FIRST + i
    first, last, back = blocks[(team, day)]
    vals = [team, day, f'=COUNT(B{first}:B{last})', f'=COUNTIF(J{first}:J{last},"Eigen lead*")',
            f'=COUNTIF(J{first}:J{last},"*was Emiel + Bas*")', f'=SUM(G{first}:G{back})', f'=C{r}*20', 30,
            f'=H{back}', f'=IF(I{r}<=TIME(17,30,0),"JA","NEE")']
    for c, v in enumerate(vals, 1):
        ws.cell(r, c, v)
    ws.cell(r, 9).number_format = HHMM
    ws.row_dimensions[r].height = 24
tot = OV_FIRST + len(order)
ws.cell(tot, 1, 'Totaal').font = BOLD
for c in 'CDEFG':
    ws[f'{c}{tot}'] = f'=SUM({c}{OV_FIRST}:{c}{tot-1})'
    ws[f'{c}{tot}'].font = BOLD
rng = f'J{OV_FIRST}:J{tot-1}'
ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"NEE"'], fill=RED_FILL, font=Font(bold=True, color='FF9C0006')))
ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"JA"'], fill=GREEN_FILL, font=Font(color='FF006100')))
ws.conditional_formatting.add(f'I{OV_FIRST}:I{tot-1}', late)

# ---- herverdeling ----
ROUTES_END = row - 1  # last row of the day blocks; lookups stop here
bar(row, 'HERVERDELING | leads die eerst bij Emiel + Bas stonden (andere leadowner)', WHITE_B12, TITLE_FILL)
row += 1
header(row, ['Nieuw team', 'Nieuwe dag', 'Lead', 'Leadowner', 'Oorspronkelijk', 'Plaats', 'Stop (nieuw)',
             'Aankomst', None, 'Waarom hier'])
row += 1
WHY = {
    'Rotterdam': 'Past in de Den Haag / Rotterdam-route van Friso + Thierry',
    'Heinenoord': 'Net onder Rotterdam, tussen Ahoy en Nieuwerkerk',
    'Zevenaar': 'Ligt ca. 30 min van Nijmegen, dus eerste stop van de Nijmegen / Brabant-route',
    'Zwanenburg': 'Op de weg naar Woodstock (Bloemendaal)',
}
for team, day, r in sorted(new_rows.values(), key=lambda x: (list(TEAMS).index(x[0]), x[1], x[2])):
    lead = ws.cell(r, 3).value
    d = leads[lead]
    why = WHY.get(d['place'])
    if not why:
        why = {('Yasmijn + Steven', 1): 'Op de terugweg vanuit Soest (A1), vlak bij de A10',
               ('Friso + Thierry', 2): 'Past in de Noord / grachtengordel-route van dag 2',
               ('Friso + Thierry', 3): 'Past in de Westpoort / Sloterdijk / West-route van dag 3'}[(team, day)]
    vals = [team, day, lead, d['owner'], f"Emiel + Bas, dag {d['day']}", d['place'], f'=B{r}', f'=H{r}', None, why]
    for c, v in enumerate(vals, 1):
        ws.cell(row, c, v).alignment = WRAP
    ws.cell(row, 8).number_format = HHMM
    ws.cell(row, 3).fill = NEW_FILL
    row += 1
row += 1

# ---- aandachtspunten ----
bar(row, 'AANDACHTSPUNTEN', WHITE_B12, TITLE_FILL)
row += 1
notes = [
    ('Friso + Thierry, dag 1', 'Terug ca. 17:35, dus ca. 5 min na 17:30. Oplossing: om 09:15 vertrekken, of The Zoo (Heinenoord) '
                               'apart inplannen (dan terug ca. 16:50). De heenrit van 93 min naar Scheveningen is al ruim geschat.'),
    ('Yasmijn + Steven, dag 2', 'Terug ca. 17:56, dus ca. 26 min na 17:30, door Listen to Flow (Zevenaar). Oplossing: om 09:00 vertrekken '
                                '(terug ca. 17:26), of Listen to Flow apart / later inplannen (dan terug ca. 16:58).'),
    ('Emiel + Bas', "Doen alleen hun eigen leads (Sheet1): Toekomst Muziek, Torpedotheater, Bijlmerbios, Theater 't Kapelletje, "
                    'Bloemencorso en Supper Club. Die plannen ze zelf en staan niet in deze planning.'),
    ('The Zoo', "Postcode staat als '274 LH'. Waarschijnlijk is dat 3274 LH (Heinenoord); graag controleren."),
    ('Listen to Flow / Jaarbeurs', "Adres staat nu wel in de route. In het tabblad 'Adressen controleren' staan ze nog als ontbrekend; "
                                   'ingang en contactpersoon vooraf bevestigen.'),
    ('Yasmijn + Steven, dag 1', 'Lunch staat volgens de afspraak (4 stops, dus op de terugreis) aan het eind van de dag. In de praktijk '
                                'lunch in Groningen; de terugtijd verandert daar niet door.'),
    ('Rijtijden', 'Bestaande trajecten hebben de rijtijden uit de oorspronkelijke planning gehouden. Nieuwe trajecten zijn handmatig '
                  'geschat (geen live Maps); parkeren en spits kunnen afwijken.'),
]
for label, text in notes:
    ws.cell(row, 1, label).font = BOLD
    ws.cell(row, 1).alignment = WRAP
    ws.cell(row, 2, text)
    row += 1

ws.freeze_panes = None

# ---- leadoverzicht: lead + leadowner, team en dag opgehaald uit 'Planning 2 teams' ----
ov = wb.create_sheet('Leadoverzicht', 0)
ov.sheet_properties.tabColor = '17243A'
for col, w in zip('ABCDEFGH', [6, 30, 14, 20, 8, 4, 14, 10]):
    ov.column_dimensions[col].width = w
for c in range(1, 6):
    ov.cell(1, c).font, ov.cell(1, c).fill = WHITE_B14, TITLE_FILL
ov.cell(1, 1, 'ADE Door-to-Door | Leadoverzicht P-Logic')
ov.row_dimensions[1].height = 18.95
ov.cell(2, 1, "Team en Dag komen automatisch uit het tabblad 'Planning 2 teams'. "
              'Poppodium Volt is weggelaten (dubbel met Volt).')
FIRST = 4
for c, v in enumerate(['#', 'Lead', 'Leadowner', 'Team', 'Dag'], 1):
    cell = ov.cell(FIRST, c, v)
    cell.font, cell.fill = WHITE_B11, HEAD_FILL
P = "'Planning 2 teams'!"
r = FIRST + 1
for team in TEAMS:
    for day in (1, 2, 3):
        for lead, _ in ROUTES[(team, day)][0]:
            match = f'MATCH(B{r},{P}$C$1:$C${ROUTES_END},0)'
            ov.cell(r, 1, r - FIRST)
            ov.cell(r, 2, lead)
            ov.cell(r, 3, leads[lead]['owner'])
            ov.cell(r, 4, f'=IFERROR(INDEX({P}$A$1:$A${ROUTES_END},{match}),"Niet ingepland")')
            ov.cell(r, 5, f'=IFERROR(COUNTIF({P}$A$1:INDEX({P}$A$1:$A${ROUTES_END},{match}),"DAG *"),"")')
            r += 1
LAST = r - 1
OWNER_FILLS = [('Yasmijn', 'FFD9EAD3'), ('Friso', 'FFF4CCCC'), ('Thierry', 'FFCFE2F3'), ('Odin', 'FFFCE5CD')]
for name, color in OWNER_FILLS:
    ov.conditional_formatting.add(f'C{FIRST+1}:C{LAST}', CellIsRule(
        operator='equal', formula=[f'"{name}"'], fill=PatternFill('solid', fgColor=color)))
for c, v in enumerate(['Legenda', 'Leads'], 7):
    cell = ov.cell(FIRST, c, v)
    cell.font, cell.fill = WHITE_B11, HEAD_FILL
for i, (name, color) in enumerate(OWNER_FILLS, FIRST + 1):
    ov.cell(i, 7, name).fill = PatternFill('solid', fgColor=color)
    ov.cell(i, 8, f'=COUNTIF($C${FIRST+1}:$C${LAST},G{i})')
tot = FIRST + 1 + len(OWNER_FILLS)
ov.cell(tot, 7, 'Totaal').font = BOLD
ov.cell(tot, 8, f'=SUM(H{FIRST+1}:H{tot-1})').font = BOLD
ov.freeze_panes = f'A{FIRST+1}'

wb.active = 0
for s in wb.worksheets:
    s.sheet_view.tabSelected = (s.title == 'Leadoverzicht')
wb.save(dst)
print('saved', dst, 'rows', row, 'blocks', blocks)
